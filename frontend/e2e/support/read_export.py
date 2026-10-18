"""Lectores independientes de los exporters de producción."""
import csv
import json
import sys
from datetime import datetime
from pathlib import Path


path = Path(sys.argv[1])
raw = path.read_bytes()
if path.suffix == ".csv":
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream, delimiter=";"))
    result = dict(format="csv", bom=raw.startswith(b"\xef\xbb\xbf"), headers=rows[0], rows=rows[1:])
elif path.suffix == ".xlsx":
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Reportes"]
    rows = list(sheet.iter_rows(values_only=True))
    result = dict(format="xlsx", sheets=workbook.sheetnames, headers=rows[0], rows=rows[1:],
                  dateCells=[isinstance(row[7], datetime) for row in rows[1:]],
                  dateFormats=[sheet.cell(index, 8).number_format for index in range(2, sheet.max_row + 1)])
    workbook.close()
elif path.suffix == ".pdf":
    from pypdf import PdfReader
    reader = PdfReader(path, strict=True)
    result = dict(format="pdf", signature=raw[:4].decode("ascii"), pages=len(reader.pages),
                  title=reader.metadata.title, texts=[page.extract_text() for page in reader.pages],
                  pageSizes=[[float(page.mediabox.width), float(page.mediabox.height)] for page in reader.pages])
    # Render opcional para QA visual; la validación del PDF siempre usa el parser real.
    try:
        import pypdfium2 as pdfium
    except ImportError:
        result["rendered"] = []
    else:
        document = pdfium.PdfDocument(str(path))
        result["rendered"] = []
        for index in range(len(document)):
            image = path.with_name(f"{path.stem}-page-{index + 1}.png")
            document[index].render(scale=1.5).to_pil().save(image)
            result["rendered"].append(str(image))
        document.close()
else:
    raise ValueError("Formato no admitido")
# ASCII en el transporte JSON evita depender de la página de códigos de Windows;
# los archivos se leen y validan con su codificación real, incluidos los acentos.
print(json.dumps(result, ensure_ascii=True, default=str))
