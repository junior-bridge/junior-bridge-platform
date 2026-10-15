import csv
import re
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO, StringIO
from unittest.mock import patch

from bson import ObjectId
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.postulations.models import Postulation
from apps.projects.models import Project
from apps.users.models import User
from .exporters import HEADERS


class ReportExportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        def user(role, name):
            return User.objects.create_user(email=f"{name}@example.test", password="Local-test-912!",
                name=name, surname="Local", role=role)
        cls.owner = user(User.Role.CLIENT, "owner")
        cls.tester = user(User.Role.TESTER, "tester")
        cls.admin = user(User.Role.ADMIN, "admin")
        cls.stranger = user(User.Role.TESTER, "stranger")
        cls.other_client = user(User.Role.CLIENT, "other-client")
        cls.project = Project.objects.create(title="Proyecto ñ", description="Prueba", repository="https://example.test",
            demo_url="https://example.test", technologies="Django", modality="REMOTE", client=cls.owner)
        cls.postulation = Postulation.objects.create(id_project=cls.project, id_tester=cls.tester, status="accepted")

    def setUp(self):
        self.url = reverse("report-export")
        self.client.force_authenticate(user=self.tester)
        self.documents = []
        for index, (severity, state, title) in enumerate([
            ("low", "PENDING", 'Login con ñ, "comillas"'),
            ("medium", "REVIEW", "Búsqueda"),
            ("high", "CLOSED", "Formulario"),
            ("critical", None, "Carrito"),
        ]):
            report = dict(_id=ObjectId(), id_postulation=self.postulation.pk, title=title,
                description='Descripción completa, con "comillas"; ñ\nSegunda línea <b>literal</b>',
                severity=severity, steps_to_reproduce=["Abrir aplicación", "Comprobar información"],
                created_at=datetime(2026, 10, 15-index, 10, 20, tzinfo=timezone.utc),
                evidence_images=["private/secret.png"], capture_evidence="https://private.test/image")
            if state:
                report.update(state=state, updated_at=datetime(2026, 10, 15, 11, tzinfo=timezone.utc))
            self.documents.append(report)
        self.documents.append(dict(self.documents[0], _id=ObjectId(), title="AJENO", id_postulation=-1))

        def find(query):
            ids = query.get("id_postulation", {}).get("$in")
            return [deepcopy(doc) for doc in self.documents if ids is None or doc["id_postulation"] in ids]
        self.find_patch = patch("apps.reports.services.reports_collection.find", side_effect=find)
        self.find = self.find_patch.start()
        self.addCleanup(self.find_patch.stop)

    def export(self, **params):
        return self.client.get(self.url, {"file_format": "csv", **params})

    def csv_rows(self, response):
        return list(csv.reader(StringIO(response.content.decode("utf-8-sig")), delimiter=";"))

    def test_csv_is_real_complete_utf8_and_private(self):
        # Arrange
        # Act
        response = self.export()
        # Assert
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        rows = self.csv_rows(response)
        self.assertEqual(rows[0], list(HEADERS))
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[1][4], self.documents[0]["description"])
        self.assertEqual(rows[1][2], "Proyecto ñ")
        self.assertIn("1. Abrir aplicación\n2. Comprobar información", rows[1])
        self.assertEqual(rows[1][10], "2")
        self.assertNotIn(b"private/secret.png", response.content)
        self.assertNotIn(b"private.test", response.content)
        self.assertNotIn(b"AJENO", response.content)
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        self.assertEqual(response["X-Report-Count"], "4")
        self.assertEqual(int(response["Content-Length"]), len(response.content))
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")

    def test_csv_neutralizes_formulas_without_losing_original_text(self):
        # Arrange
        # Act
        for formula in ('=HYPERLINK("evil")', '  +SUM(1,2)', '-1+1', '@cmd', '\t=1', '\r=1', '\n=1'):
            with self.subTest(formula=formula):
                self.documents[0]["title"] = formula
        # Assert
                self.assertEqual(self.csv_rows(self.export())[1][3], "'" + formula)

    def test_xlsx_is_real_and_user_text_cannot_become_formula(self):
        # Arrange
        # Act
        self.documents[0]["title"] = "=1+1"
        response = self.export(file_format="xlsx")
        # Assert
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Reportes"])
        sheet = workbook["Reportes"]
        self.assertEqual(sheet.max_row, 5)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["D2"].value, "=1+1")
        self.assertEqual(sheet["D2"].data_type, "s")
        self.assertEqual(sheet["E2"].value, self.documents[0]["description"])
        self.assertEqual(sheet["H2"].value, datetime(2026, 10, 15, 10, 20))
        self.assertTrue(sheet["H2"].is_date)
        self.assertTrue(sheet["E2"].alignment.wrap_text)
        self.assertEqual(sheet["K2"].value, 2)
        self.assertEqual(sheet["G5"].value, "ABIERTO")
        self.assertIsNone(sheet["I5"].value)

    def test_excel_never_silently_truncates_long_text(self):
        # Arrange
        # Act
        self.documents[0]["description"] = "ñ" * 33000
        response = self.export(file_format="xlsx")
        # Assert
        self.assertEqual(response.status_code, 400)
        self.assertIn("Excel", response.data["detail"])
        self.assertNotIn("Content-Disposition", response)

    def test_pdf_is_real_multipage_with_long_text_and_escaped_markup(self):
        # Arrange
        # Act
        self.documents[0]["description"] = '<img src="https://private.test"> & ñ\n' * 500
        response = self.export(file_format="pdf")
        # Assert
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", response.content[-1024:])
        self.assertGreater(len(re.findall(rb"/Type\s*/Page\b", response.content)), 1)

    def test_anonymous_cannot_export(self):
        # Arrange
        # Act
        self.client.force_authenticate(user=None)
        # Assert
        self.assertEqual(self.export().status_code, 401)
        self.find.assert_not_called()

    def test_authorization_is_same_as_list_for_every_role(self):
        # Arrange
        # Act
        for user, count in [(self.admin, 5), (self.tester, 4), (self.owner, 4), (self.stranger, 0), (self.other_client, 0)]:
            with self.subTest(role=user.email):
                self.client.force_authenticate(user=user)
                listing = self.client.get(reverse("report-list"))
        # Assert
                self.assertEqual(len(listing.data), count)
                response = self.export()
                if count:
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(len(self.csv_rows(response))-1, count)
                    self.assertEqual({row[0] for row in self.csv_rows(response)[1:]}, {row["_id"] for row in listing.data})
                else:
                    self.assertEqual(response.status_code, 400)
                    self.assertNotIn("Content-Disposition", response)

    def test_filters_match_frontend_fields_and_utc_day(self):
        # Arrange
        # Act
        for params, count in [
            ({"q": "LOGIN"}, 1), ({"q": "media"}, 1), ({"q": "REVISIÓN"}, 1),
            ({"q": "abierto"}, 2), ({"q": "critical"}, 1), ({"q": str(self.postulation.pk)}, 4),
            ({"q": "  Login"}, 0), ({"date": "2026-10-14"}, 1),
            ({"date": "2026-10-14", "q": "búsqueda"}, 1),
            ({"date": "2026-10-14", "q": "login"}, 0), ({"q": "no existe"}, 0),
        ]:
            with self.subTest(params=params):
                response = self.export(**params)
        # Assert
                self.assertEqual(response.status_code, 200 if count else 400)
                if count:
                    self.assertEqual(len(self.csv_rows(response))-1, count)

    def test_filters_apply_to_all_formats(self):
        # Arrange
        # Act
        for fmt in ("csv", "pdf", "xlsx"):
            with self.subTest(fmt=fmt):
                response = self.export(file_format=fmt, q="Login", date="2026-10-15")
        # Assert
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["X-Report-Count"], "1")

    def test_invalid_parameters_do_not_expand_access(self):
        # Arrange
        # Act
        for params in ({"file_format": "exe"}, {"date": "2026-02-30"}, {"date": "2026-1-1"},
                {"user_id": self.admin.pk}, {"postulation_id": -1}, {"q": ["a", "b"]}):
            with self.subTest(params=params):
                response = self.export(**params)
        # Assert
                self.assertEqual(response.status_code, 400)
                self.assertNotIn("Content-Disposition", response)
        self.assertEqual(self.client.get(self.url).status_code, 400)

    def test_no_partial_success_when_generation_fails(self):
        # Arrange
        # Act
        with patch("apps.reports.views.generate_csv", side_effect=RuntimeError("generation failed")):
        # Assert
            with self.assertLogs("apps.reports.views", level="ERROR"):
                response = self.export()
        self.assertEqual(response.status_code, 500)
        self.assertNotIn("Content-Disposition", response)
        self.assertIn("archivo completo", response.data["detail"])

    def test_export_limit_returns_error_without_truncating(self):
        # Arrange
        # Act
        with patch("apps.reports.views.MAX_EXPORT_REPORTS", 2):
            response = self.export()
        # Assert
        self.assertEqual(response.status_code, 400)
        self.assertNotIn("Content-Disposition", response)

    def test_missing_fields_and_orphan_project_are_explicit(self):
        # Arrange
        # Act
        self.client.force_authenticate(user=self.admin)
        response = self.export()
        rows = self.csv_rows(response)
        # Assert
        self.assertEqual(rows[-1][2], "")
        self.assertEqual(rows[4][6], "ABIERTO")
        self.assertEqual(rows[4][8], "")
