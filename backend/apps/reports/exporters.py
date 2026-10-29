"""Generate complete files in memory before sending a successful response."""

import csv
import math
from datetime import datetime, timezone
from functools import lru_cache
from io import BytesIO, StringIO
from pathlib import Path
from xml.sax.saxutils import escape

import reportlab
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from apps.postulations.models import Postulation
from .services import SEVERITY_LABELS, STATUS_LABELS


HEADERS = (
    "ID del reporte", "ID de postulación", "Proyecto", "Título", "Descripción",
    "Severidad", "Estado", "Fecha de creación (UTC)", "Fecha de actualización (UTC)",
    "Pasos para reproducir", "Cantidad de evidencias",
)
CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


class ExportValidationError(ValueError):
    """A user-facing, safe error for a format's documented limits."""


def _date(value):
    if not value:
        return None
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)


def export_rows(reports):
    projects = dict(Postulation.objects.filter(
        pk__in={report["id_postulation"] for report in reports}
    ).values_list("pk", "id_project__title"))
    return [(
        report["_id"], report["id_postulation"], projects.get(report["id_postulation"], ""),
        report["title"], report["description"],
        SEVERITY_LABELS.get(report["severity"], report["severity"]),
        STATUS_LABELS.get(report.get("state") if report.get("state") is not None else "PENDING", report.get("state") or ""),
        _date(report.get("created_at")), _date(report.get("updated_at")),
        "\n".join(f"{i}. {step}" for i, step in enumerate(report.get("steps_to_reproduce", []), 1)),
        len(report.get("evidence_images") or []) + bool(report.get("capture_evidence")),
    ) for report in reports]


def _text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M:%S")
    return str(value)


def _csv_text(value):
    text = _text(value)
    # Spreadsheet apps may ignore leading whitespace before a formula marker.
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def generate_csv(rows):
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(HEADERS)
    writer.writerows([_csv_text(value) for value in row] for row in rows)
    return output.getvalue().encode("utf-8-sig")


def generate_xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reportes"
    sheet.append(HEADERS)
    widths = [28, 18, 30, 36, 65, 14, 18, 25, 25, 55, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2D6A4F")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for row_index, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(row_index, column)
            if isinstance(value, datetime):
                cell.value = value.replace(tzinfo=None)  # UTC, stated in the header.
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            elif isinstance(value, str):
                if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
                    raise ExportValidationError("Un texto excede los límites de Excel. Exportá como CSV o PDF.")
                cell.value = value
                cell.data_type = "s"  # Never interpret user input as a formula.
            else:
                cell.value = value
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        lines = max(sum(max(1, math.ceil(len(line) / (width - 2))) for line in _text(value).split("\n"))
            for value, width in zip(row, widths))
        # Excel's maximum row height is 409 points; cell values remain complete.
        sheet.row_dimensions[row_index].height = min(409, max(30, lines * 15))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@lru_cache(maxsize=1)
def _register_pdf_fonts():
    font_dir = Path(reportlab.__file__).parent / "fonts"
    pdfmetrics.registerFont(TTFont("JBVera", str(font_dir / "Vera.ttf")))
    pdfmetrics.registerFont(TTFont("JBVeraBold", str(font_dir / "VeraBd.ttf")))


def generate_pdf(rows, *, query, date, generated_at):
    _register_pdf_fonts()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4),
        leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36,
        title="Reportes de Bugs - JuniorBridge", author="JuniorBridge")
    body = ParagraphStyle("Body", fontName="JBVera", fontSize=9, leading=13, spaceAfter=8, splitLongWords=True)
    heading = ParagraphStyle("Heading", parent=body, fontName="JBVeraBold", fontSize=16, leading=22, textColor=colors.HexColor("#2d6a4f"))
    small = ParagraphStyle("Small", parent=body, fontSize=8, leading=11, spaceAfter=0)
    header = ParagraphStyle("Header", parent=small, fontName="JBVeraBold", textColor=colors.white)

    def paragraph(value, style=body):
        return Paragraph(escape(_text(value)).replace("\n", "<br/>"), style)

    story = [paragraph("Reportes de Bugs - JuniorBridge", heading),
        paragraph(f"Generado: {_text(generated_at)} UTC | Registros: {len(rows)}", small),
        paragraph(f"Búsqueda: {query or 'Sin filtro'} | Fecha (UTC): {date or 'Sin filtro'}", small),
        Spacer(1, 16)]
    summary_headers = ["Reporte / Postulación", "Proyecto", "Título", "Severidad", "Estado", "Creación (UTC)", "Evidencias"]
    summary = [[paragraph(value, header) for value in summary_headers]]
    for row in rows:
        summary.append([paragraph(value, small) for value in (
            f"{row[0]}\nPostulación: {row[1]}", row[2], row[3], row[5], row[6], row[7], row[10],
        )])
    table = LongTable(summary, colWidths=[125, 105, 220, 65, 80, 110, 64], repeatRows=1, splitInRow=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2d6a4f")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f0eb")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#98410f")),
    ]))
    story.extend([table, Spacer(1, 20), paragraph("Detalle completo de los reportes", heading)])
    for row in rows:
        story.append(paragraph(f"{row[3]} | ID: {row[0]}", ParagraphStyle("Report", parent=body, fontName="JBVeraBold", keepWithNext=True)))
        for label, value in (
            ("Postulación", row[1]), ("Proyecto", row[2]), ("Severidad", row[5]),
            ("Estado", row[6]), ("Creación (UTC)", row[7]), ("Actualización (UTC)", row[8]),
            ("Cantidad de evidencias", row[10]), ("Descripción", row[4]), ("Pasos para reproducir", row[9]),
        ):
            story.append(paragraph(f"{label}: {_text(value) if value is not None and value != '' else 'Sin dato'}"))
        story.append(Spacer(1, 12))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("JBVera", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(36, 20, "JuniorBridge | Fechas en UTC | Exportación de reportes autorizados")
        canvas.drawRightString(landscape(A4)[0] - 36, 20, f"Página {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
